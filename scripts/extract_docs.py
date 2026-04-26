"""Extract text from docx and pdf docs into plain text for review."""
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "obsolete"
OUT = ROOT / "scripts" / "extracted"
OUT.mkdir(parents=True, exist_ok=True)


def extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    chunks = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            chunks.append(text)
    for ti, table in enumerate(doc.tables):
        chunks.append(f"\n[TABLE {ti}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages):
        out.append(f"\n===== PAGE {i+1} =====\n")
        out.append(page.extract_text() or "")
    return "\n".join(out)


for name in [
    "AIDI2026_Methodology_v2.docx",
    "Lidar Data Description.docx",
]:
    src = DOCS / name
    dst = OUT / (src.stem + ".txt")
    text = extract_docx(src)
    dst.write_text(text, encoding="utf-8")
    print(f"Wrote {dst} ({len(text):,} chars)")

src = DOCS / "AIDI2026_Research_Proposal_v2.pdf"
dst = OUT / (src.stem + ".txt")
text = extract_pdf(src)
dst.write_text(text, encoding="utf-8")
print(f"Wrote {dst} ({len(text):,} chars)")
